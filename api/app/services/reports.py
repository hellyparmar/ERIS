"""Downloadable reports (CSV and Excel).

Each report is a function returning rows (list of dicts) for a period and outlet scope; COLUMNS fixes the
column order and the human-readable headers. Excel files get a second sheet with the report's provenance
(period, outlets, generation time, data source) so an exported file can be understood on its own.
"""
from __future__ import annotations

import csv
import io
from collections.abc import Callable
from dataclasses import dataclass

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app import clock
from app.models import Category, DatasetInfo, InventoryItem, Outlet, Product, Sale, SaleItem
from app.services import analytics as A


@dataclass
class Report:
    key: str
    title: str
    description: str
    columns: list[tuple[str, str]]  # (field, header)
    build: Callable[[Session, A.DateRange, list[int] | None], list[dict]]
    uses_period: bool = True
    min_role: str = "viewer"


def _sales_daily(db: Session, rng: A.DateRange, outlet_ids: list[int] | None) -> list[dict]:
    names = A.outlet_names(db)
    rows = db.execute(select(Sale.sale_date, Sale.outlet_id, func.count(Sale.id), func.sum(Sale.total),
                             func.sum(Sale.items_count), func.sum(Sale.discount), func.sum(Sale.tax_amount)).where(
        *A._sale_filters(rng, outlet_ids)).group_by(Sale.sale_date, Sale.outlet_id).order_by(
        Sale.sale_date, Sale.outlet_id)).all()
    return [{"date": d.isoformat(), "outlet": names.get(o, o), "bills": n, "revenue": round(rev, 2),
             "avg_bill": round(rev / n, 2) if n else 0, "units": round(units or 0, 1), "discount": round(disc or 0, 2),
             "tax": round(tax or 0, 2), "net_sales": round(rev - (tax or 0), 2)}
            for d, o, n, rev, units, disc, tax in rows]


def _outlets(db: Session, rng: A.DateRange, outlet_ids: list[int] | None) -> list[dict]:
    return A.outlet_performance(db, rng, outlet_ids)


def _products(db: Session, rng: A.DateRange, outlet_ids: list[int] | None) -> list[dict]:
    return A.abc_analysis(db, rng, outlet_ids)["products"]


def _categories(db: Session, rng: A.DateRange, outlet_ids: list[int] | None) -> list[dict]:
    return A.category_breakdown(db, rng, outlet_ids)


def _gst(db: Session, rng: A.DateRange, outlet_ids: list[int] | None) -> list[dict]:
    """Output tax by outlet state, HSN code and rate. Retail (B2C) sales are intra-state supplies, so the tax
    splits equally into CGST and SGST."""
    q = A._item_query(rng, outlet_ids, Outlet.state, Product.hsn_code, Product.tax_rate,
                      func.sum(SaleItem.line_total), func.sum(SaleItem.tax_amount), func.sum(SaleItem.quantity)).join(
        Product, Product.id == SaleItem.product_id).join(Outlet, Outlet.id == SaleItem.outlet_id).group_by(
        Outlet.state, Product.hsn_code, Product.tax_rate).order_by(Outlet.state, Product.hsn_code)
    out = []
    for state, hsn, rate, gross, tax, qty in db.execute(q).all():
        tax = tax or 0.0
        out.append({"state": state, "hsn_code": hsn, "tax_rate": rate, "quantity": round(qty, 1),
                    "taxable_value": round(gross - tax, 2), "cgst": round(tax / 2, 2), "sgst": round(tax / 2, 2),
                    "total_tax": round(tax, 2), "invoice_value": round(gross, 2)})
    return out


def _inventory(db: Session, _rng: A.DateRange, outlet_ids: list[int] | None) -> list[dict]:
    q = select(InventoryItem, Product, Outlet, Category.name).join(Product, Product.id == InventoryItem.product_id).join(
        Outlet, Outlet.id == InventoryItem.outlet_id).join(Category, Category.id == Product.category_id).where(
        Product.is_active.is_(True)).order_by(Outlet.id, Product.sku)
    if outlet_ids:
        q = q.where(InventoryItem.outlet_id.in_(outlet_ids))
    out = []
    for inv, p, o, cat in db.execute(q).all():
        qty = inv.quantity
        status = "out of stock" if qty <= 0 else ("low" if qty <= inv.reorder_level else "ok")
        out.append({"outlet": o.name, "sku": p.sku, "product": p.name, "category": cat, "unit": p.unit,
                    "quantity": qty, "reorder_level": inv.reorder_level, "status": status,
                    "unit_cost": p.cost_price, "stock_value_cost": round(max(qty, 0) * p.cost_price, 2),
                    "stock_value_retail": round(max(qty, 0) * p.selling_price, 2)})
    return out


def _reorder(db: Session, _rng: A.DateRange, outlet_ids: list[int] | None) -> list[dict]:
    from app.services.forecasting import reorder_suggestions

    return reorder_suggestions(db, outlet_ids)


def _customers(db: Session, _rng: A.DateRange, outlet_ids: list[int] | None) -> list[dict]:
    return A.customer_segments(db, outlet_ids)["segments"]


def _anomalies(db: Session, rng: A.DateRange, outlet_ids: list[int] | None) -> list[dict]:
    from app.services import anomalies as AN

    rows = AN.detect(db, rng.start, rng.end, outlet_ids)
    for a in rows:
        a["explanation"] = AN.explain(db, a)
    return rows


REPORTS: dict[str, Report] = {r.key: r for r in [
    Report("sales-daily", "Daily sales by outlet", "Bills, revenue, average bill, units, discounts and tax per day.",
           [("date", "Date"), ("outlet", "Outlet"), ("bills", "Bills"), ("revenue", "Revenue (incl. tax)"),
            ("avg_bill", "Average bill"), ("units", "Units"), ("discount", "Discount"), ("tax", "Tax"),
            ("net_sales", "Net sales (excl. tax)")], _sales_daily),
    Report("outlets", "Outlet performance", "Revenue, bills, profit and growth per outlet versus the previous period.",
           [("code", "Code"), ("name", "Outlet"), ("city", "City"), ("revenue", "Revenue"), ("orders", "Bills"),
            ("avg_basket", "Average bill"), ("units", "Units"), ("profit", "Gross profit"), ("margin_pct", "Margin %"),
            ("share_pct", "Share %"), ("previous_revenue", "Previous period revenue"), ("change_pct", "Change %")],
           _outlets),
    Report("products", "Product performance (ABC)", "Revenue, units and margin per product with ABC class.",
           [("sku", "SKU"), ("name", "Product"), ("category", "Category"), ("units", "Units"), ("revenue", "Revenue"),
            ("profit", "Gross profit"), ("margin_pct", "Margin %"), ("share_pct", "Share %"),
            ("cumulative_pct", "Cumulative %"), ("class", "ABC class")], _products),
    Report("categories", "Category performance", "Revenue, units, profit and growth per category.",
           [("category", "Category"), ("revenue", "Revenue"), ("units", "Units"), ("profit", "Gross profit"),
            ("margin_pct", "Margin %"), ("share_pct", "Share %"), ("change_pct", "Change vs previous %")],
           _categories),
    Report("gst-summary", "GST output tax summary (demo)",
           "Taxable value and CGST/SGST by state, HSN code and rate. For analysis only - not for tax filing.",
           [("state", "State"), ("hsn_code", "HSN"), ("tax_rate", "GST rate %"), ("quantity", "Quantity"),
            ("taxable_value", "Taxable value"), ("cgst", "CGST"), ("sgst", "SGST"), ("total_tax", "Total tax"),
            ("invoice_value", "Invoice value")], _gst, min_role="manager"),
    Report("inventory", "Inventory valuation", "Current stock per outlet and product at cost and retail value.",
           [("outlet", "Outlet"), ("sku", "SKU"), ("product", "Product"), ("category", "Category"), ("unit", "Unit"),
            ("quantity", "Quantity"), ("reorder_level", "Reorder level"), ("status", "Status"),
            ("unit_cost", "Unit cost"), ("stock_value_cost", "Value at cost"),
            ("stock_value_retail", "Value at selling price")], _inventory, uses_period=False),
    Report("reorder", "Reorder recommendations", "Demand-based order quantities for items running low.",
           [("outlet", "Outlet"), ("sku", "SKU"), ("product", "Product"), ("supplier", "Supplier"),
            ("current_stock", "Stock"), ("on_order", "On order"), ("avg_daily_demand", "Daily demand"),
            ("days_of_cover", "Days of cover"), ("lead_time_days", "Lead time (days)"),
            ("suggested_qty", "Suggested qty"), ("estimated_cost", "Estimated cost"), ("urgency", "Urgency")],
           _reorder, uses_period=False),
    Report("customer-segments", "Customer segments (RFM)", "Customers grouped by recency, frequency and spend.",
           [("segment", "Segment"), ("description", "Description"), ("customers", "Customers"),
            ("revenue", "Revenue"), ("revenue_share_pct", "Revenue share %"), ("avg_recency_days", "Avg days since visit"),
            ("avg_orders", "Avg bills"), ("avg_spend", "Avg spend")], _customers, uses_period=False),
    Report("anomalies", "Unusual days", "Outlet-days with unusual revenue and the evidence found for each.",
           [("day", "Date"), ("outlet", "Outlet"), ("direction", "Direction"), ("actual", "Actual revenue"),
            ("expected", "Expected revenue"), ("change_pct", "Change %"), ("z_score", "Robust z"),
            ("severity", "Severity"), ("explained_by", "Context"), ("explanation", "Evidence")], _anomalies),
]}

ROLE_RANK = {"viewer": 0, "staff": 0, "manager": 1, "admin": 2}


def available(role: str) -> list[Report]:
    return [r for r in REPORTS.values() if ROLE_RANK.get(role, 0) >= ROLE_RANK[r.min_role]]


def build(db: Session, report: Report, rng: A.DateRange, outlet_ids: list[int] | None) -> list[list]:
    rows = report.build(db, rng, outlet_ids)
    return [[r.get(k) for k, _ in report.columns] for r in rows]


def provenance(db: Session, report: Report, rng: A.DateRange, outlet_ids: list[int] | None) -> list[tuple[str, str]]:
    names = A.outlet_names(db)
    ds = db.scalar(select(DatasetInfo).order_by(DatasetInfo.id.desc()))
    sources = sorted(s for (s,) in db.execute(select(Sale.source).distinct()).all())
    info = [("Report", report.title), ("Description", report.description),
            ("Period", f"{rng.start.isoformat()} to {rng.end.isoformat()}" if report.uses_period else "Current"),
            ("Outlets", ", ".join(names.get(o, str(o)) for o in outlet_ids) if outlet_ids else "All outlets"),
            ("Generated at", clock.now().strftime("%Y-%m-%d %H:%M")),
            ("Sales sources in database", ", ".join(sources) or "none")]
    if ds:
        info.append(("Dataset", f"{ds.data_source} (generator {ds.generator_version}, seed {ds.random_seed}, "
                                f"{ds.period_start} to {ds.period_end})"))
    return info


def to_csv(report: Report, rows: list[list]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h for _, h in report.columns])
    w.writerows(rows)
    return buf.getvalue()


def to_xlsx(report: Report, rows: list[list], info: list[tuple[str, str]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31].replace("/", "-")
    ws.append([h for _, h in report.columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, (key, header) in enumerate(report.columns, start=1):
        width = max([len(header)] + [len(str(r[i - 1])) for r in rows[:500] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
        if any(t in key for t in ("revenue", "value", "cost", "profit", "bill", "discount", "tax", "sales", "spend",
                                  "gst", "actual", "expected")) and "pct" not in key:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = "#,##0.00"
    meta = wb.create_sheet("About this report")
    for k, v in info:
        meta.append([k, v])
    meta.column_dimensions["A"].width = 28
    meta.column_dimensions["B"].width = 90
    for cell in meta["A"]:
        cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
