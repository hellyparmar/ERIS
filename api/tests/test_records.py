"""Reports and exports, demo GST invoices and the audit log."""
import io

from app.db import SessionLocal
from app.models import Customer, Sale
from app.services.gst import is_valid_gstin
from app.services.invoices import amount_in_words


def test_reports_list_preview_and_export(client, admin, viewer):
    keys = {r["key"] for r in client.get("/api/reports", headers=admin).json()}
    assert {"sales-daily", "outlets", "products", "gst-summary", "inventory", "reorder"} <= keys
    assert "gst-summary" not in {r["key"] for r in client.get("/api/reports", headers=viewer).json()}
    assert client.get("/api/reports/gst-summary", headers=viewer).status_code == 403

    from datetime import date, timedelta

    start, end = date.today() - timedelta(days=36), date.today() - timedelta(days=30)  # inside the seeded history
    p = client.get(f"/api/reports/sales-daily?start={start}&end={end}", headers=admin).json()
    # one row per outlet-day with sales (an injected closure day has none)
    assert 7 * 4 <= p["total_rows"] <= 7 * 5 and len(p["rows"][0]) == len(p["columns"])
    csv = client.get("/api/reports/outlets/export?format=csv&period=30d", headers=admin)
    assert csv.status_code == 200 and csv.text.startswith("Code,Outlet") and csv.text.count("\n") >= 6

    from openpyxl import load_workbook

    x = client.get("/api/reports/gst-summary/export?format=xlsx&period=30d", headers=admin)
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames[1] == "About this report"
    ws = wb.worksheets[0]
    assert ws["A1"].value == "State" and ws.max_row > 1
    # CGST and SGST split the tax equally for counter sales
    for row in ws.iter_rows(min_row=2, values_only=True):
        assert abs(row[5] - row[6]) < 0.02


def test_reports_respect_outlet_scope(client, manager):
    rows = client.get("/api/reports/inventory?limit=500", headers=manager).json()["rows"]
    assert rows and {r[0] for r in rows} == {"Andheri West"}
    assert client.get("/api/reports/inventory?outlet_id=2", headers=manager).status_code == 403


def _sale_id(outlet_id: int) -> int:
    with SessionLocal() as db:
        return db.query(Sale).filter_by(outlet_id=outlet_id, status="completed").order_by(Sale.id.desc()).first().id


def test_invoice_intra_state_and_pdf(client, admin):
    sale_id = _sale_id(1)
    r = client.post("/api/invoices", headers=admin, json={"sale_id": sale_id})
    assert r.status_code == 201, r.text
    inv = r.json()
    assert inv["created"] and inv["is_demo"] and inv["supply_type"] == "intra_state"
    assert inv["number"].startswith("MUMAND/") and len(inv["number"]) <= 16
    assert is_valid_gstin(inv["seller_gstin"]) and inv["seller_gstin"].startswith("27")
    assert inv["igst"] == 0 and abs(inv["cgst"] - inv["sgst"]) < 0.02
    assert abs(inv["taxable_value"] + inv["cgst"] + inv["sgst"] - inv["total"]) < 0.02
    # issuing again returns the same invoice
    again = client.post("/api/invoices", headers=admin, json={"sale_id": sale_id}).json()
    assert again["id"] == inv["id"] and not again["created"]
    pdf = client.get(f"/api/invoices/{inv['id']}/pdf", headers=admin)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert any(i["id"] == inv["id"] for i in client.get("/api/invoices", headers=admin).json()["items"])


def test_invoice_inter_state_uses_igst(client, admin):
    sale_id = _sale_id(3)  # Bengaluru outlet (Karnataka, 29)
    bad = client.post("/api/invoices", headers=admin, json={"sale_id": sale_id, "buyer_gstin": "27AAAAA0000A1Z0"})
    assert bad.status_code == 422
    r = client.post("/api/invoices", headers=admin, json={"sale_id": sale_id, "buyer_gstin": "27AAPFU0939F1ZV"})
    inv = r.json()
    assert inv["supply_type"] == "inter_state" and inv["place_of_supply"] == "27"
    assert inv["cgst"] == 0 and inv["sgst"] == 0 and inv["igst"] > 0
    assert inv["seller_gstin"].startswith("29")


def test_invoice_access_rules(client, admin, manager, viewer):
    other = _sale_id(2)
    assert client.post("/api/invoices", headers=manager, json={"sale_id": other}).status_code == 403
    assert client.post("/api/invoices", headers=viewer, json={"sale_id": other}).status_code == 403
    inv = client.post("/api/invoices", headers=admin, json={"sale_id": other}).json()
    assert client.get(f"/api/invoices/{inv['id']}", headers=manager).status_code == 403


def test_amount_in_words():
    assert amount_in_words(123456.5) == "Rupees One Lakh Twenty Three Thousand Four Hundred Fifty Six and Fifty Paise Only"
    assert amount_in_words(0) == "Rupees Zero Only"


def test_audit_log_records_changes(client, admin, manager):
    r = client.post("/api/customers", headers=admin, json={"name": "Audit Probe", "phone": "9123400001"})
    assert r.status_code == 201, r.text
    cid = r.json()["id"]
    client.put(f"/api/customers/{cid}", headers=admin, json={"name": "Audit Probe 2", "phone": "9123400001"})
    items = client.get("/api/audit?entity=customer", headers=admin).json()["items"]
    mine = [a for a in items if a["entity_id"] == str(cid)]
    assert {a["action"] for a in mine} == {"customer.create", "customer.update"}
    upd = next(a for a in mine if a["action"] == "customer.update")
    assert upd["details"]["name"] == {"from": "Audit Probe", "to": "Audit Probe 2"} and upd["user"].startswith("Aditi")
    assert client.get("/api/audit", headers=manager).status_code == 403
    with SessionLocal() as db:
        assert db.query(Customer).filter_by(id=cid).one().name == "Audit Probe 2"


def test_record_times_use_the_business_clock(client, admin):
    """Audit entries, stock movements and runs are stamped in business time (Asia/Kolkata), like sales."""
    from datetime import datetime

    from app import clock

    client.post("/api/inventory/adjust", headers=admin, json={"outlet_id": 1, "product_id": 7, "mode": "add", "quantity": 1})
    for path, key in (("/api/audit", "at"), ("/api/inventory/movements", "created_at")):
        rows = client.get(path, headers=admin).json()
        rows = rows["items"] if isinstance(rows, dict) else rows
        newest = max(datetime.fromisoformat(r[key]) for r in rows)
        assert abs((newest - clock.now()).total_seconds()) < 120, (path, newest, clock.now())
