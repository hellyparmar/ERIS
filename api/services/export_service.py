"""
Export Service - Multi-format Report Generation
Supports PDF, Excel, and CSV exports for analytics and data
"""

import io
import csv
from datetime import datetime
from typing import List, Dict, Any, Optional
import logging

logger = logging.getLogger(__name__)

# Try importing optional dependencies
try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False
    logger.warning("ReportLab not installed. PDF export will be unavailable.")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning("openpyxl not installed. Excel export will be unavailable.")


class ExportService:
    """Service for exporting data in multiple formats"""
    
    def __init__(self):
        self.styles = getSampleStyleSheet() if REPORTLAB_AVAILABLE else None
    
    # ==================== CSV EXPORT ====================
    
    def export_to_csv(self, data: List[Dict[str, Any]], filename: str = "export.csv") -> bytes:
        """
        Export data to CSV format
        
        Args:
            data: List of dictionaries containing the data
            filename: Name for the file
            
        Returns:
            CSV file as bytes
        """
        if not data:
            return b""
        
        output = io.StringIO()
        
        # Get headers from first row
        headers = list(data[0].keys())
        
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerows(data)
        
        return output.getvalue().encode('utf-8')
    
    # ==================== EXCEL EXPORT ====================
    
    def export_to_excel(
        self, 
        data: List[Dict[str, Any]], 
        sheet_name: str = "Data",
        title: str = "Export Report"
    ) -> bytes:
        """
        Export data to Excel format with styling
        
        Args:
            data: List of dictionaries containing the data
            sheet_name: Name of the worksheet
            title: Title for the report
            
        Returns:
            Excel file as bytes
        """
        if not OPENPYXL_AVAILABLE:
            raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")
        
        if not data:
            return b""
        
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        # Add title
        ws.merge_cells('A1:' + get_column_letter(len(data[0])) + '1')
        title_cell = ws['A1']
        title_cell.value = title
        title_cell.font = Font(size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30
        
        # Add timestamp
        ws.merge_cells('A2:' + get_column_letter(len(data[0])) + '2')
        timestamp_cell = ws['A2']
        timestamp_cell.value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        timestamp_cell.font = Font(size=10, italic=True)
        timestamp_cell.alignment = Alignment(horizontal="center")
        
        # Add headers
        headers = list(data[0].keys())
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num)
            cell.value = header.replace('_', ' ').title()
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Add data
        for row_num, row_data in enumerate(data, 5):
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=row_num, column=col_num)
                value = row_data.get(header, "")
                cell.value = value
                
                # Alternate row colors
                if row_num % 2 == 0:
                    cell.fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        
        # Auto-adjust column widths
        for col_num, header in enumerate(headers, 1):
            max_length = len(header)
            for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=col_num, max_col=col_num):
                try:
                    if len(str(row[0].value)) > max_length:
                        max_length = len(str(row[0].value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[get_column_letter(col_num)].width = adjusted_width
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output.getvalue()
    
    # ==================== PDF EXPORT ====================
    
    def export_to_pdf(
        self,
        data: List[Dict[str, Any]],
        title: str = "Export Report",
        subtitle: str = None
    ) -> bytes:
        """
        Export data to PDF format with professional styling
        
        Args:
            data: List of dictionaries containing the data
            title: Title for the report
            subtitle: Optional subtitle
            
        Returns:
            PDF file as bytes
        """
        if not REPORTLAB_AVAILABLE:
            raise ImportError("ReportLab is required for PDF export. Install with: pip install reportlab")
        
        if not data:
            return b""
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elements = []
        
        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1e3a8a'),
            spaceAfter=12,
            alignment=1  # Center
        )
        elements.append(Paragraph(title, title_style))
        
        # Subtitle
        if subtitle:
            subtitle_style = ParagraphStyle(
                'CustomSubtitle',
                parent=self.styles['Normal'],
                fontSize=12,
                textColor=colors.grey,
                spaceAfter=12,
                alignment=1
            )
            elements.append(Paragraph(subtitle, subtitle_style))
        
        # Timestamp
        timestamp = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        timestamp_style = ParagraphStyle(
            'Timestamp',
            parent=self.styles['Normal'],
            fontSize=9,
            textColor=colors.grey,
            spaceAfter=20,
            alignment=1
        )
        elements.append(Paragraph(timestamp, timestamp_style))
        elements.append(Spacer(1, 0.2*inch))
        
        # Table
        headers = list(data[0].keys())
        table_data = [
            [header.replace('_', ' ').title() for header in headers]
        ]
        
        for row in data:
            table_data.append([str(row.get(header, "")) for header in headers])
        
        # Create table
        table = Table(table_data)
        
        # Style the table
        table.setStyle(TableStyle([
            # Header styling
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3b82f6')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            
            # Data styling
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
            
            # Grid
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('BOX', (0, 0), (-1, -1), 1, colors.black),
        ]))
        
        elements.append(table)
        
        # Build PDF
        doc.build(elements)
        buffer.seek(0)
        
        return buffer.getvalue()
    
    # ==================== SPECIALIZED EXPORTS ====================
    
    def export_sales_report(
        self,
        sales_data: List[Dict[str, Any]],
        format: str = "pdf",
        date_range: str = None
    ) -> bytes:
        """Export sales report in specified format"""
        
        title = "Sales Report"
        subtitle = f"Period: {date_range}" if date_range else None
        
        if format.lower() == "pdf":
            return self.export_to_pdf(sales_data, title, subtitle)
        elif format.lower() == "excel":
            return self.export_to_excel(sales_data, "Sales Data", title)
        elif format.lower() == "csv":
            return self.export_to_csv(sales_data)
        else:
            raise ValueError(f"Unsupported format: {format}")
    
    def export_inventory_report(
        self,
        inventory_data: List[Dict[str, Any]],
        format: str = "pdf"
    ) -> bytes:
        """Export inventory report in specified format"""
        
        title = "Inventory Report"
        subtitle = f"Total Items: {len(inventory_data)}"
        
        if format.lower() == "pdf":
            return self.export_to_pdf(inventory_data, title, subtitle)
        elif format.lower() == "excel":
            return self.export_to_excel(inventory_data, "Inventory", title)
        elif format.lower() == "csv":
            return self.export_to_csv(inventory_data)
        else:
            raise ValueError(f"Unsupported format: {format}")
    
    def export_invoice_report(
        self,
        invoice_data: List[Dict[str, Any]],
        format: str = "pdf"
    ) -> bytes:
        """Export invoice report in specified format"""
        
        title = "Invoice Report"
        total_amount = sum(float(inv.get('amount', 0)) for inv in invoice_data)
        subtitle = f"Total Invoices: {len(invoice_data)} | Total Amount: ₹{total_amount:,.2f}"
        
        if format.lower() == "pdf":
            return self.export_to_pdf(invoice_data, title, subtitle)
        elif format.lower() == "excel":
            return self.export_to_excel(invoice_data, "Invoices", title)
        elif format.lower() == "csv":
            return self.export_to_csv(invoice_data)
        else:
            raise ValueError(f"Unsupported format: {format}")
